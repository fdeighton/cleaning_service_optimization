"""One unified Parker workbook: Schedule Grid + Clean Schedule, nothing else.

Output: output/clean schedule/Parker_Clean_Schedule.xlsx
  Sheet 1 'Schedule Grid'  — Day x Time Block grid of responsibility labels
  Sheet 2 'Clean Schedule' — one row per assignment (the detailed schedule)

Reuses the existing, already-verified builders so content matches the standalone
versions exactly.
"""
from __future__ import annotations

import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_area_inventory import classify_area, clean_location  # noqa: E402
from build_schedule_grid import NOTE, WEEK, build_grid, cell_label, resolve_source_csv  # noqa: E402
from build_clean_schedule import (  # noqa: E402
    HEADERS as CS_HEADERS, WIDTHS as CS_WIDTHS, WRAP_COLS as CS_WRAP,
    clean, day_key, time_key,
)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "output", "clean schedule", "Parker_Clean_Schedule.xlsx")

# shared styling
HEAD_FILL = PatternFill("solid", fgColor="1A1A1A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BAND = PatternFill("solid", fgColor="F4F4F4")
STRIPE = PatternFill("solid", fgColor="F6F6F6")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_grid_sheet(wb: Workbook) -> tuple[list, int, int]:
    employees, grid, review, cells = build_grid()
    ws = wb.create_sheet("Schedule Grid", 0)
    ws.sheet_view.showGridLines = False
    ncols = 2 + len(employees)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    n = ws.cell(1, 1, NOTE)
    n.font = Font(italic=True, size=9, color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 26

    headers = ["Day", "Time Block"] + employees
    for j, h in enumerate(headers, start=1):
        c = ws.cell(2, j, h); c.fill = HEAD_FILL; c.font = HEAD_FONT; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[2].height = 20

    day_order = {d: i for i, d in enumerate(WEEK)}
    for i, row in enumerate(grid):
        r = i + 3
        shade = day_order[row[0]] % 2 == 1
        for j, val in enumerate(row, start=1):
            c = ws.cell(r, j, val); c.border = BORDER; c.alignment = WRAP
            if j == 1:
                c.font = Font(bold=True)
            if shade:
                c.fill = BAND

    for j, w in enumerate([13, 18] + [20] * len(employees), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{len(grid) + 2}"

    _write_legend(ws, cells, ncols)
    return employees, len(grid), len(review)


# Friendly descriptions for the cluster labels (legend on the Schedule Grid).
LEGEND_DESC = {
    "Lobby Cluster": "Ground-floor common-area sweep — lobby, enterphone, security, café, mail, elevators, washrooms",
    "Outdoor / Patio": "Exterior grounds — patios, breezeway, dog run, BBQ, debris & trash pickup",
    "Parking": "Parking levels and parking-area amenities",
    "Guest Suites": "Guest and model suites",
    "Corridor Program": "Corridors, hallways and stairwells",
    "Elevator / Chute Program": "Elevators and garbage chutes",
    "Back-of-House": "Storage, housekeeping, offices, mechanical / admin rooms",
    "Washrooms": "Washrooms and change rooms",
    "Specials": "Re-check / re-sweep / detail and as-required touch-ups",
    "Amenities": "Resident amenity spaces spanning multiple floors",
    "Lunch": "Scheduled lunch break",
    "Break": "Scheduled break",
}


def _trunc(s: str, n: int = 30) -> str:
    return s if len(s) <= n else s[: n - 1] + "…"


def _write_legend(ws, cells, ncols):
    """Legend with concrete coverage per label, derived from the actual rows
    (so 'Specials' shows it spans the 5th / 38th / 39th floors, etc.)."""
    from collections import Counter, defaultdict
    label_cells = Counter()
    label_locs = defaultdict(Counter)
    for items in cells.values():
        if not items:
            continue
        lab = cell_label(items)
        label_cells[lab] += 1
        for it in items:
            label_locs[lab][it["loc"]] += 1

    def base_of(lab):
        if lab in LEGEND_DESC:
            return LEGEND_DESC[lab]
        if lab.endswith("Amenities"):
            return f"Amenity spaces on the {lab.replace(' Amenities', '').strip()}"
        return ""

    entries = []
    for lab in sorted(label_cells, key=lambda l: (-label_cells[l], l)):
        top = label_locs[lab].most_common(6)
        locs = "; ".join(_trunc(l) for l, _ in top)
        more = len(label_locs[lab]) - len(top)
        if more > 0:
            locs += f"  (+{more} more)"
        base = base_of(lab)
        entries.append((lab, (base + " — " if base else "") + "Covers: " + locs))

    lc, mc = ncols + 2, ncols + 3   # one blank column gap, then label + meaning
    ws.merge_cells(start_row=2, start_column=lc, end_row=2, end_column=mc)
    t = ws.cell(2, lc, "Legend — Cluster Labels (areas covered)")
    t.font = HEAD_FONT; t.fill = HEAD_FILL; t.alignment = Alignment(vertical="center")
    ws.cell(2, mc).fill = HEAD_FILL
    for c in (lc, mc):
        ws.cell(2, c).border = BORDER
    for k, (lab, desc) in enumerate(entries):
        r = 3 + k
        a = ws.cell(r, lc, lab); a.font = Font(bold=True); a.border = BORDER
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(r, mc, desc); b.border = BORDER; b.alignment = WRAP
        ws.row_dimensions[r].height = 30
    ws.column_dimensions[get_column_letter(lc)].width = 22
    ws.column_dimensions[get_column_letter(mc)].width = 70


def write_clean_sheet(wb: Workbook) -> int:
    with open(resolve_source_csv(), newline="", encoding="utf-8-sig") as f:
        src = list(csv.DictReader(f))
    rows = []
    for r in src:
        s, e = clean(r.get("start_time", "")), clean(r.get("end_time", ""))
        loc = clean(r.get("location", ""))
        rows.append([
            clean(r.get("employee_name", "")), clean(r.get("day_of_week", "")), s, e,
            f"{s} – {e}" if s and e else (s or e),
            loc, classify_area(clean_location(loc)),
            clean(r.get("task", "")), clean(r.get("frequency", "")), clean(r.get("notes", "")),
        ])
    rows.sort(key=lambda x: (x[0].lower(), day_key(x[1]), time_key(x[2]), x[5].lower()))

    ws = wb.create_sheet("Clean Schedule")
    ws.sheet_view.showGridLines = False
    for j, h in enumerate(CS_HEADERS, start=1):
        c = ws.cell(1, j, h); c.fill = HEAD_FILL; c.font = HEAD_FONT; c.border = BORDER
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 20
    for i, row in enumerate(rows, start=2):
        for j, val in enumerate(row, start=1):
            c = ws.cell(i, j, val); c.border = BORDER
            c.alignment = WRAP if CS_HEADERS[j - 1] in CS_WRAP else Alignment(vertical="top")
            if i % 2 == 0:
                c.fill = STRIPE
    for j, h in enumerate(CS_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = CS_WIDTHS[h]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CS_HEADERS))}{len(rows) + 1}"
    return len(rows)


def main():
    wb = Workbook()
    wb.remove(wb.active)
    employees, grid_rows, review = write_grid_sheet(wb)
    sched_rows = write_clean_sheet(wb)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    try:
        wb.save(OUT)
    except PermissionError:
        print(f"WARNING: {os.path.basename(OUT)} is open in Excel — close it and rerun.")
        return

    print("output path:", os.path.relpath(OUT, REPO))
    print("sheets:", wb.sheetnames)
    print("employees:", employees)
    print("Schedule Grid rows:", grid_rows, "| Review Needed:", review)
    print("Clean Schedule rows:", sched_rows)


if __name__ == "__main__":
    main()
