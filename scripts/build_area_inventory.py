"""Build the Area Inventory Request package for the Development team.

Reads (never modifies) the schedule CSVs in schedule_data/, discovers every
unique (property_name, location) combination, applies conservative cleaning,
classification, and prioritization, and writes a business-ready request package
to output/.

Strictly an inventory-request task: no square footage is estimated, no workload
or analytics is produced.
"""
from __future__ import annotations

import glob
import os
import re

import pandas as pd

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCHEDULE_DIR = os.path.join(REPO, "schedule_data")
OUTPUT_DIR = os.path.join(REPO, "output")
CONFIG_DIR = os.path.join(REPO, "config")
PROPERTY_NAME_MAP_CSV = os.path.join(CONFIG_DIR, "property_name_map.csv")

# Columns only Development can complete — always emitted blank.
DEV_BLANK_COLUMNS = ["square_footage", "development_notes"]
REQUEST_COLUMNS = [
    "property_name", "raw_location", "proposed_clean_location",
    "proposed_area_type", "priority", *DEV_BLANK_COLUMNS,
]


# ---------------------------------------------------------------------------
# Load
# ---------------------------------------------------------------------------
def load_schedules() -> pd.DataFrame:
    files = sorted(glob.glob(os.path.join(SCHEDULE_DIR, "*_cleaning_responsibilities.csv")))
    if not files:
        raise FileNotFoundError(f"No schedule CSVs found in {SCHEDULE_DIR}")
    frames = []
    for f in files:
        df = pd.read_csv(f, dtype=str, keep_default_na=False)
        frames.append(df)
    df = pd.concat(frames, ignore_index=True)
    # The source uses `property`; the project's canonical name is property_name.
    df = df.rename(columns={"property": "property_name"})
    df["property_name"] = df["property_name"].str.strip().map(property_name_map())
    df["location"] = df["location"].str.strip()
    return df


def property_name_map() -> "dict[str, str]":
    """raw property label -> canonical property_name. Unmapped labels pass through.

    Resolves the Elm/Ledbury labelling: 'Elm', 'Ledbury', and 'Elm/Ledbury' all
    collapse to the single combined property 'Elm/Ledbury'.
    """
    m: dict[str, str] = {}
    if os.path.exists(PROPERTY_NAME_MAP_CSV):
        mp = pd.read_csv(PROPERTY_NAME_MAP_CSV, dtype=str, keep_default_na=False)
        m = {r["raw_property"].strip(): r["property_name"].strip() for _, r in mp.iterrows()}

    class _PassThrough(dict):
        def __missing__(self, key):  # any label not in the map stays as-is
            return key

    return _PassThrough(m)


# ---------------------------------------------------------------------------
# Phase 2 — conservative cleaning
# ---------------------------------------------------------------------------
_LEADING_JUNK = re.compile(r"^[\s\-–•\*•]+")
_MULTISPACE = re.compile(r"\s{2,}")


def clean_location(raw: str) -> str:
    """Apply only obvious, non-destructive normalization.

    Trims whitespace, collapses internal runs of spaces, strips leading
    bullet/dash junk, and tidies spacing around separators. Capitalization and
    wording are left intact so abbreviations (BBQ, STOA, P1) and proper nouns are
    never damaged; case-variant near-duplicates are surfaced in the manual-review
    file instead of being silently merged.
    """
    s = raw.strip()
    s = _LEADING_JUNK.sub("", s)
    s = s.replace(" ,", ",").replace(" ;", ";")
    s = _MULTISPACE.sub(" ", s)
    return s.strip()


# ---------------------------------------------------------------------------
# Phase 3 — conservative area classification
# ---------------------------------------------------------------------------
# (keyword, area_type) — first match wins, so order = precedence.
AREA_RULES: list[tuple[str, str]] = [
    # explicit ambiguous / surface-only references -> Other / Unknown
    ("building-wide", "Other / Unknown"),
    ("not specified", "Other / Unknown"),
    ("all mopped", "Other / Unknown"),
    ("all staff", "Other / Unknown"),
    # pet amenities before generic "spa"
    ("dog spa", "Amenity"),
    ("pet spa", "Amenity"),
    ("dog wash", "Amenity"),
    # washrooms / change rooms
    ("washroom", "Washroom / Change Room"),
    ("restroom", "Washroom / Change Room"),
    ("change room", "Washroom / Change Room"),
    ("changeroom", "Washroom / Change Room"),
    ("shower", "Washroom / Change Room"),
    # pool / spa
    ("pool", "Pool / Spa"),
    ("lido", "Pool / Spa"),
    ("sauna", "Pool / Spa"),
    ("hot tub", "Pool / Spa"),
    ("jacuzzi", "Pool / Spa"),
    ("steam room", "Pool / Spa"),
    # fitness
    ("gym", "Fitness"),
    ("fitness", "Fitness"),
    ("yoga", "Fitness"),
    ("the temple", "Fitness"),
    # vertical transport / circulation
    ("elevator", "Elevator"),
    ("stairwell", "Stairwell"),
    ("staircase", "Stairwell"),
    ("stair well", "Stairwell"),
    # loading dock before generic corridor/bay
    ("loading dock", "Loading Dock"),
    ("loading bay", "Loading Dock"),
    # parking before lobby (e.g. "Parking Lobbies")
    ("parking", "Parking"),
    ("garage", "Parking"),
    ("parkade", "Parking"),
    # waste / garbage (chutes are garbage chutes)
    ("garbage", "Waste / Garbage"),
    ("chute", "Waste / Garbage"),
    ("trash", "Waste / Garbage"),
    ("waste", "Waste / Garbage"),
    ("compactor", "Waste / Garbage"),
    ("recycl", "Waste / Garbage"),
    # corridor / hallway (exterior corridor stays a corridor)
    ("corridor", "Corridor / Hallway"),
    ("hallway", "Corridor / Hallway"),
    ("breezeway", "Corridor / Hallway"),
    ("exit", "Corridor / Hallway"),
    # lobby / entrance
    ("lobby", "Lobby"),
    ("entrance", "Lobby"),
    ("enterphone", "Lobby"),
    ("foyer", "Lobby"),
    ("vestibule", "Lobby"),
    ("reception", "Lobby"),
    # terrace / outdoor
    ("terrace", "Terrace / Outdoor"),
    ("patio", "Terrace / Outdoor"),
    ("rooftop", "Terrace / Outdoor"),
    ("balcony", "Terrace / Outdoor"),
    ("courtyard", "Terrace / Outdoor"),
    ("dog run", "Terrace / Outdoor"),
    ("dog exit", "Terrace / Outdoor"),
    ("driveway", "Terrace / Outdoor"),
    ("barbeque", "Terrace / Outdoor"),
    ("barbecue", "Terrace / Outdoor"),
    ("bbq", "Terrace / Outdoor"),
    ("exterior", "Terrace / Outdoor"),
    ("outdoor", "Terrace / Outdoor"),
    # guest / model suites
    ("guest suite", "Guest Suite"),
    ("model suite", "Guest Suite"),
    ("model unit", "Guest Suite"),
    ("suite #", "Guest Suite"),
    # office / admin
    ("mngt", "Office / Admin"),
    ("management", "Office / Admin"),
    ("office", "Office / Admin"),
    ("admin", "Office / Admin"),
    ("security", "Office / Admin"),
    ("list office", "Office / Admin"),
    ("mail room", "Office / Admin"),
    ("mailroom", "Office / Admin"),
    ("parcel", "Office / Admin"),
    # storage / housekeeping
    ("storage", "Storage"),
    ("housekeeping", "Storage"),
    ("janitor", "Storage"),
    ("supply room", "Storage"),
    ("locker", "Storage"),
    # mechanical / back-of-house
    ("mechanical", "Mechanical / Back-of-House"),
    ("electrical", "Mechanical / Back-of-House"),
    ("boiler", "Mechanical / Back-of-House"),
    ("sprinkler", "Mechanical / Back-of-House"),
    ("maintenance", "Mechanical / Back-of-House"),
    ("utility", "Mechanical / Back-of-House"),
    ("telecom", "Mechanical / Back-of-House"),
    # amenities (broad resident amenity bucket, lower precedence)
    ("amenit", "Amenity"),
    ("lounge", "Amenity"),
    ("games room", "Amenity"),
    ("game room", "Amenity"),
    ("party room", "Amenity"),
    ("entertainment", "Amenity"),
    ("kitchen", "Amenity"),
    ("cafe", "Amenity"),
    ("dining", "Amenity"),
    ("library", "Amenity"),
    ("theatre", "Amenity"),
    ("theater", "Amenity"),
    ("media room", "Amenity"),
    ("bowling", "Amenity"),
    ("golf", "Amenity"),
    ("basketball", "Amenity"),
    ("court", "Amenity"),
    ("studio", "Amenity"),
    ("simulator", "Amenity"),
    ("sport", "Amenity"),
    ("kids", "Amenity"),
    ("adventure", "Amenity"),
    ("green house", "Amenity"),
    ("greenhouse", "Amenity"),
    ("pizza oven", "Amenity"),
    ("clinic", "Amenity"),
    ("meeting room", "Amenity"),
    ("board room", "Amenity"),
    ("co-work", "Amenity"),
    ("coworking", "Amenity"),
    ("music", "Amenity"),
    ("juice", "Amenity"),
    ("laundry", "Amenity"),
    ("bar", "Amenity"),
    ("spa", "Amenity"),
]


def classify_area(clean_loc: str) -> str:
    text = clean_loc.lower()
    # A name that *starts* with "stair(s)" is a stairwell; mid-string "w/stairs"
    # (inside lounge/gym descriptions) must not trigger this, hence startswith.
    if text.startswith("stair"):
        return "Stairwell"
    for kw, area_type in AREA_RULES:
        if kw in text:
            return area_type
    return "Other / Unknown"


# ---------------------------------------------------------------------------
# Phase 4 — prioritization
# ---------------------------------------------------------------------------
MAJOR_COMMON = {"Lobby", "Corridor / Hallway", "Elevator", "Amenity", "Fitness", "Pool / Spa"}
# Surface/feature references that likely don't need their own SF row.
MAY_NOT_NEED_SF = (
    "plaque", "trash can", "fire cabinet", "threshold", "window", "glass",
    "carpet", "baseboard", "ventilation", "wallpaper", "sign", "metal",
)


def assign_priority(area_type: str, occ_count: int, has_daily: bool, clean_loc: str) -> str:
    low_text = clean_loc.lower()
    # Surface/feature references (carpets, baseboards, glass, thresholds...) are
    # not discrete areas Development can size — always Low, even if frequent.
    if any(t in low_text for t in MAY_NOT_NEED_SF):
        return "Low Priority"
    # High: appears many times (top ~10%, occ>=8), a recurring major common area,
    # or a recurring daily location.
    if (
        occ_count >= 8
        or (area_type in MAJOR_COMMON and occ_count >= 2)
        or (has_daily and occ_count >= 4)
    ):
        return "High Priority"
    # Low: rare, unclear, or administrative references.
    if (
        occ_count <= 1
        or area_type == "Other / Unknown"
        or area_type == "Office / Admin"
    ):
        return "Low Priority"
    return "Medium Priority"


# ---------------------------------------------------------------------------
# Phase 7 — manual-review flagging
# ---------------------------------------------------------------------------
# Markers of genuinely ambiguous / non-discrete references. Deliberately excludes
# "&", "and", "w/" — those appear in many valid descriptive area names.
AMBIGUOUS_MARKERS = (
    "building-wide", "not specified", "all mopped", "all staff", "(all",
    "various", " etc",
)


def build_inventory(df: pd.DataFrame) -> pd.DataFrame:
    # Phase 1 — unique (property, location) with occurrence count + frequency.
    grouped = (
        df.groupby(["property_name", "location"])
        .agg(
            occurrence_count=("location", "size"),
            frequencies=("frequency", lambda s: sorted(set(s.str.strip()))),
        )
        .reset_index()
        .rename(columns={"location": "raw_location"})
    )

    grouped["proposed_clean_location"] = grouped["raw_location"].apply(clean_location)
    grouped["proposed_area_type"] = grouped["proposed_clean_location"].apply(classify_area)
    grouped["has_daily"] = grouped["frequencies"].apply(
        lambda fs: any(f.lower() == "daily" for f in fs)
    )
    grouped["priority"] = grouped.apply(
        lambda r: assign_priority(
            r["proposed_area_type"], r["occurrence_count"], r["has_daily"],
            r["proposed_clean_location"],
        ),
        axis=1,
    )

    # Duplicate detection: same property + same case-insensitive clean name but
    # >1 distinct raw spelling => possible duplicate.
    grouped["_clean_key"] = (
        grouped["property_name"].str.lower() + "||" +
        grouped["proposed_clean_location"].str.lower()
    )
    dup_keys = (
        grouped.groupby("_clean_key")["raw_location"].nunique()
    )
    dup_keys = set(dup_keys[dup_keys > 1].index)

    def review_reasons(r) -> list[str]:
        reasons = []
        if r["proposed_area_type"] == "Other / Unknown":
            reasons.append("area type unclear")
        if r["_clean_key"] in dup_keys:
            reasons.append("possible duplicate (case/spelling variant)")
        low = r["proposed_clean_location"].lower()
        if any(m in low for m in AMBIGUOUS_MARKERS):
            reasons.append("ambiguous / bundled naming")
        if len(r["proposed_clean_location"]) <= 3:
            reasons.append("very short / unclear name")
        return reasons

    grouped["_review_reasons"] = grouped.apply(review_reasons, axis=1)
    grouped["needs_review"] = grouped["_review_reasons"].apply(bool)
    return grouped


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------
_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def _sanitize_sheet_name(name: str, used: set[str]) -> str:
    """Make an Excel-legal, unique sheet name (no :\\/?*[], <=31 chars)."""
    s = _INVALID_SHEET_CHARS.sub("-", name).strip()[:31] or "Sheet"
    base = s
    i = 1
    while s.lower() in used:
        suffix = f"_{i}"
        s = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(s.lower())
    return s


def _style_sheet(ws, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + df[col].astype(str).str.len().tolist()) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 60)
    ws.freeze_panes = "A2"


def _write_workbook(xlsx_path: str, sheets: list[tuple[str, pd.DataFrame]]) -> None:
    """Write one xlsx with multiple styled sheets: [(sheet_name, df), ...]."""
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xl:
        for name, d in sheets:
            d.to_excel(xl, index=False, sheet_name=name)
            _style_sheet(xl.sheets[name], d)


def write_table(df: pd.DataFrame, basename: str, sheet_name: str) -> None:
    """Single-sheet table: CSV + one-sheet xlsx."""
    csv_path = os.path.join(OUTPUT_DIR, f"{basename}.csv")
    xlsx_path = os.path.join(OUTPUT_DIR, f"{basename}.xlsx")
    df.to_csv(csv_path, index=False)
    used: set[str] = set()
    try:
        _write_workbook(xlsx_path, [(_sanitize_sheet_name(sheet_name, used), df)])
    except PermissionError:
        print(f"WARNING: could not write {basename}.xlsx (file open/locked). "
              f"{basename}.csv was updated. Close the file and rerun.")


def write_request(df: pd.DataFrame, basename: str = "area_inventory_request") -> None:
    """Request package: a flat global CSV + an xlsx with a global sheet AND one
    sheet per property."""
    csv_path = os.path.join(OUTPUT_DIR, f"{basename}.csv")
    xlsx_path = os.path.join(OUTPUT_DIR, f"{basename}.xlsx")
    df.to_csv(csv_path, index=False)  # CSV stays a single global table

    used: set[str] = set()
    sheets: list[tuple[str, pd.DataFrame]] = [
        (_sanitize_sheet_name("All Properties", used), df)
    ]
    for prop in sorted(df["property_name"].unique()):
        sub = df[df["property_name"] == prop].reset_index(drop=True)
        sheets.append((_sanitize_sheet_name(prop, used), sub))
    try:
        _write_workbook(xlsx_path, sheets)
    except PermissionError:
        print(f"WARNING: could not write {basename}.xlsx (file open/locked). "
              f"{basename}.csv was updated. Close the file and rerun.")


def main() -> None:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    df = load_schedules()
    inv = build_inventory(df)

    # Phase 5 — request file (Development-facing columns; blanks left empty).
    request = inv.sort_values(
        ["property_name", "priority", "proposed_clean_location"],
        key=lambda s: s.map({"High Priority": 0, "Medium Priority": 1, "Low Priority": 2}).fillna(s),
    ).copy()
    for c in DEV_BLANK_COLUMNS:
        request[c] = ""
    request_out = request[REQUEST_COLUMNS]
    write_request(request_out)

    # Phase 6 — per-property summary.
    def prio_count(g, p):
        return int((g["priority"] == p).sum())

    summary_rows = []
    for prop, g in inv.groupby("property_name"):
        summary_rows.append({
            "property_name": prop,
            "unique_location_count": int(g["raw_location"].nunique()),
            "unique_area_type_count": int(g["proposed_area_type"].nunique()),
            "high_priority_count": prio_count(g, "High Priority"),
            "medium_priority_count": prio_count(g, "Medium Priority"),
            "low_priority_count": prio_count(g, "Low Priority"),
            "unknown_area_count": int((g["proposed_area_type"] == "Other / Unknown").sum()),
        })
    summary = pd.DataFrame(summary_rows).sort_values("property_name", ignore_index=True)
    write_table(summary, "area_inventory_summary", "Summary")

    # Phase 7 — manual review file (CSV only, as specified).
    review = inv[inv["needs_review"]].copy()
    review["review_reasons"] = review["_review_reasons"].apply(lambda rs: "; ".join(rs))
    review = review[[
        "property_name", "raw_location", "proposed_clean_location",
        "proposed_area_type", "priority", "occurrence_count", "review_reasons",
    ]].sort_values(["property_name", "proposed_clean_location"], ignore_index=True)
    review.to_csv(os.path.join(OUTPUT_DIR, "manual_review_locations.csv"), index=False)

    # Console validation report (Phase 8).
    print("properties_discovered:", inv["property_name"].nunique())
    print("total_location_references (property x location rows):", len(inv))
    print("total_unique_location_strings:", df["location"].nunique())
    print("\narea_type_distribution:")
    print(inv["proposed_area_type"].value_counts().to_string())
    print("\npriority_distribution:")
    print(inv["priority"].value_counts().to_string())
    print("\nlocations_requiring_review:", int(inv["needs_review"].sum()))
    print("\noutput files written to:", OUTPUT_DIR)


if __name__ == "__main__":
    main()
