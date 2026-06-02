"""Build Parker_Operations_Workbook.xlsx — a clean, presentation-ready operational
workbook from Parker_cleaning_responsibilities.csv.

Audience: Operations, Development, Asset Management, Leadership. The workbook does
NOT reproduce the WhiteRose source layout; it re-presents the data so a first-time
reader can understand it. Schedule-data only — no effort, durations, or SF.

Reuses the approved area-type classification and review logic from
build_area_inventory.py so categories stay consistent across deliverables.
"""
from __future__ import annotations

import os
import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_area_inventory import (  # noqa: E402
    AMBIGUOUS_MARKERS, classify_area, clean_location,
)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_CSV = os.path.join(REPO, "schedule_data", "Parker_cleaning_responsibilities.csv")
OUTPUT_XLSX = os.path.join(REPO, "Parker_Operations_Workbook.xlsx")

WEEK = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
WEEK_FULL = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
ABBR = {"mon": "Mon", "tue": "Tue", "tues": "Tue", "wed": "Wed", "thu": "Thu",
        "thur": "Thu", "thurs": "Thu", "fri": "Fri", "sat": "Sat", "sun": "Sun"}
FULL_TO_ABBR = dict(zip([d.lower() for d in WEEK_FULL], WEEK))


# ---------------------------------------------------------------------------
# Time / day helpers
# ---------------------------------------------------------------------------
_TIME_RE = re.compile(r"^\s*(\d{1,2}):(\d{2})\s*([AaPp][Mm])?\s*$")


def to_minutes(t: str):
    m = _TIME_RE.match(str(t or ""))
    if not m:
        return None
    h = int(m.group(1)) % 12
    if (m.group(3) or "").lower() == "pm":
        h += 12
    return h * 60 + int(m.group(2))


def fmt_minutes(mins) -> str:
    if mins is None or pd.isna(mins):
        return ""
    mins = int(mins)
    h, mm = divmod(mins, 60)
    ap = "AM" if h < 12 else "PM"
    hr = h % 12 or 12
    return f"{hr}:{mm:02d} {ap}"


def expand_days(value: str) -> list[str]:
    raw = (value or "").strip()
    low = raw.lower()
    if low in FULL_TO_ABBR:
        return [FULL_TO_ABBR[low]]
    if "-" in raw:
        a, _, b = low.partition("-")
        sa, sb = ABBR.get(a.strip()), ABBR.get(b.strip())
        if sa and sb:
            i, j = WEEK.index(sa), WEEK.index(sb)
            return WEEK[i:j + 1] if i <= j else WEEK[i:] + WEEK[:j + 1]
    if low in ABBR:
        return [ABBR[low]]
    return []


def compress_days(days: set[str]) -> str:
    """Render a set of abbreviated days as compact ranges, e.g. 'Mon–Fri, Sun'."""
    idx = sorted(WEEK.index(d) for d in days if d in WEEK)
    if not idx:
        return ""
    parts, start, prev = [], idx[0], idx[0]
    for k in idx[1:]:
        if k == prev + 1:
            prev = k
            continue
        parts.append((start, prev))
        start = prev = k
    parts.append((start, prev))
    out = []
    for a, b in parts:
        out.append(WEEK[a] if a == b else f"{WEEK[a]}–{WEEK[b]}")
    return ", ".join(out)


def day_sort_key(value: str) -> int:
    days = expand_days(value)
    return WEEK.index(days[0]) if days else 99


# ---------------------------------------------------------------------------
# Load + enrich (Parker only)
# ---------------------------------------------------------------------------
def load_parker() -> pd.DataFrame:
    df = pd.read_csv(SOURCE_CSV, dtype=str, keep_default_na=False)
    for c in ["property", "employee_name", "day_of_week", "start_time",
              "end_time", "location", "task", "frequency"]:
        df[c] = df[c].str.strip()
    df["area_type"] = df["location"].apply(lambda s: classify_area(clean_location(s)))
    df["start_min"] = df["start_time"].apply(to_minutes)
    df["end_min"] = df["end_time"].apply(to_minutes)
    return df


def needs_review(location: str, area_type: str) -> bool:
    low = location.lower()
    return (
        area_type == "Other / Unknown"
        or any(m in low for m in AMBIGUOUS_MARKERS)
        or len(location.strip()) <= 3
    )


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------
def sheet_employee_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for emp, g in df.groupby("employee_name"):
        days = set()
        for v in g["day_of_week"]:
            days.update(expand_days(v))
        blocks = g[["day_of_week", "start_time", "end_time"]].drop_duplicates()
        starts = g["start_min"].dropna()
        ends = g["end_min"].dropna()
        area_counts = g["area_type"].value_counts()
        primary = ", ".join(area_counts.head(3).index.tolist())
        unknown_freq = int((g["frequency"].str.lower() == "unknown").sum())
        note = f"{unknown_freq} assignment(s) with Unknown frequency" if unknown_freq else ""
        rows.append({
            "Employee Name": emp,
            "Days Worked": compress_days(days),
            "Shift Start": fmt_minutes(starts.min() if len(starts) else None),
            "Shift End": fmt_minutes(ends.max() if len(ends) else None),
            "Total Schedule Blocks": len(blocks),
            "Distinct Locations": g["location"].nunique(),
            "Distinct Tasks": g["task"].nunique(),
            "Primary Area Types": primary,
            "Notes": note,
        })
    return pd.DataFrame(rows).sort_values("Employee Name", ignore_index=True)


def sheet_employee_schedule(df: pd.DataFrame) -> pd.DataFrame:
    s = df.copy()
    s["_day"] = s["day_of_week"].apply(day_sort_key)
    s["_start"] = s["start_min"].fillna(10**9)
    s = s.sort_values(["employee_name", "_day", "_start", "location"])
    out = s[["employee_name", "day_of_week", "start_time", "end_time",
             "location", "task", "frequency"]].copy()
    out.columns = ["Employee", "Day", "Start Time", "End Time",
                   "Location", "Task", "Frequency"]
    return out.reset_index(drop=True)


def sheet_area_coverage(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for loc, g in df.groupby("location"):
        emps = sorted(g["employee_name"].unique())
        freqs = sorted(g["frequency"].unique())
        rows.append({
            "Location": loc,
            "Area Type": g["area_type"].mode().iat[0],
            "Employees Assigned": ", ".join(emps),
            "Assignment Count": len(g),
            "Task Count": g["task"].nunique(),
            "Frequency Types": ", ".join(freqs),
        })
    return (pd.DataFrame(rows)
            .sort_values(["Assignment Count", "Location"], ascending=[False, True],
                         ignore_index=True))


def sheet_employee_area_matrix(df: pd.DataFrame) -> pd.DataFrame:
    m = (df.pivot_table(index="location", columns="employee_name",
                        aggfunc="size", fill_value=0)
         .sort_index())
    m.insert(0, "Total", m.sum(axis=1))
    m = m.sort_values("Total", ascending=False)
    m = m.reset_index().rename(columns={"location": "Location"})
    return m


def sheet_area_inventory_request(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for loc, g in df.groupby("location"):
        rows.append({
            "Property Name": g["property"].iat[0],
            "Location": loc,
            "Proposed Area Type": g["area_type"].mode().iat[0],
            "Square Footage": "",
            "Development Notes": "",
        })
    return pd.DataFrame(rows).sort_values("Location", ignore_index=True)


def sheet_schedule_insights(df: pd.DataFrame) -> pd.DataFrame:
    loc_counts = df["location"].value_counts()
    task_counts = df["task"].value_counts()
    freq_counts = df["frequency"].value_counts()
    review_locs = sorted({loc for loc, g in df.groupby("location")
                          if needs_review(loc, g["area_type"].mode().iat[0])})

    def top_list(counts, n=5):
        return "\n".join(f"{name}  ({cnt})" for name, cnt in counts.head(n).items())

    rows = [
        ("Total Employees", df["employee_name"].nunique()),
        ("Total Locations", df["location"].nunique()),
        ("Total Tasks", df["task"].nunique()),
        ("Total Schedule Assignments", len(df)),
        ("", ""),
        ("Most Serviced Locations", top_list(loc_counts)),
        ("Most Common Tasks", top_list(task_counts)),
        ("Most Common Frequencies", top_list(freq_counts)),
        ("", ""),
        ("Locations Requiring Manual Review", len(review_locs)),
        ("Review List", "\n".join(review_locs) if review_locs else "None"),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _autosize(ws, df: pd.DataFrame, wrap_cols: set[str] | None = None) -> None:
    wrap_cols = wrap_cols or set()
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        if col in wrap_cols:
            ws.column_dimensions[letter].width = 42
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            continue
        cell_lens = df[col].astype(str).str.split("\n").apply(
            lambda parts: max(len(p) for p in parts) if parts else 0)
        width = max([len(str(col))] + cell_lens.tolist()) + 2
        ws.column_dimensions[letter].width = min(width, 60)


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _add_table(ws, df: pd.DataFrame, name: str) -> None:
    ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight9", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)


def main() -> None:
    df = load_parker()

    builders = [
        ("Employee Summary", sheet_employee_summary(df), "EmployeeSummary", {"Notes"}),
        ("Employee Schedule", sheet_employee_schedule(df), "EmployeeSchedule", None),
        ("Area Coverage", sheet_area_coverage(df), "AreaCoverage",
         {"Employees Assigned", "Frequency Types"}),
        ("Employee x Area Matrix", sheet_employee_area_matrix(df), None, None),
        ("Area Inventory Request", sheet_area_inventory_request(df), "AreaInventoryRequest", None),
        ("Schedule Insights", sheet_schedule_insights(df), None, {"Value"}),
    ]

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as xl:
        for sheet_name, data, _, _ in builders:
            data.to_excel(xl, index=False, sheet_name=sheet_name)

    wb = load_workbook(OUTPUT_XLSX)
    for sheet_name, data, table_name, wrap in builders:
        ws = wb[sheet_name]
        _style_header(ws)
        _autosize(ws, data, wrap)
        ws.freeze_panes = "A2"
        if table_name:
            _add_table(ws, data, table_name)
        else:
            # Matrix / insights: header filter without full table styling.
            if sheet_name != "Schedule Insights":
                ws.auto_filter.ref = f"A1:{get_column_letter(len(data.columns))}{len(data) + 1}"
    wb.save(OUTPUT_XLSX)

    # Report
    dq = []
    unk = int((df["frequency"].str.lower() == "unknown").sum())
    if unk:
        dq.append(f"{unk} assignments have 'Unknown' frequency")
    blank_task = int(df["task"].eq("").sum())
    if blank_task:
        dq.append(f"{blank_task} rows have a blank task")
    review_n = len({loc for loc, g in df.groupby("location")
                    if needs_review(loc, g["area_type"].mode().iat[0])})
    dq.append(f"{review_n} locations flagged for manual review (ambiguous/unclassified)")
    unclass = int((df["area_type"] == "Other / Unknown").sum())
    dq.append(f"{unclass} assignments map to 'Other / Unknown' area type")

    print("workbook:", OUTPUT_XLSX)
    print("sheets:", [b[0] for b in builders])
    print("employees:", df["employee_name"].nunique())
    print("locations:", df["location"].nunique())
    print("tasks:", df["task"].nunique())
    print("assignments:", len(df))
    print("data_quality_issues:")
    for d in dq:
        print("  -", d)


if __name__ == "__main__":
    main()
