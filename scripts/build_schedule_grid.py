"""Add a human-readable 'Schedule Grid' sheet to the Parker workbook.

Rows = Day x Time Block, Columns = Employees, Cells = one concise responsibility
label summarised from that employee's assignments in that block. Labels are
derived from the actual schedule data (area-type classification); nothing is
invented, no durations/optimization. Added as the FIRST sheet of a COPY — the
original workbook and all detail sheets are left unchanged.
"""
from __future__ import annotations

import csv
import glob
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_area_inventory import classify_area, clean_location  # noqa: E402
from build_clean_schedule import clean  # noqa: E402  (same field cleanup as detail sheet)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "output", "Parker_Operations_Workbook_With_Grid.xlsx")
NOTE = ("Summary labels are derived from schedule detail records. "
        "See detail sheets for full location/task breakdown.")

WEEK = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
ABBR = {"mon": 0, "tue": 1, "tues": 1, "wed": 2, "thu": 3, "thur": 3, "thurs": 3,
        "fri": 4, "sat": 5, "sun": 6}
FULL = {d.lower(): i for i, d in enumerate(WEEK)}

AREA_LABEL = {
    "Lobby": "Lobby Cluster",
    "Parking": "Parking",
    "Guest Suite": "Guest Suites",
    "Corridor / Hallway": "Corridor Program",
    "Stairwell": "Corridor Program",
    "Elevator": "Elevator / Chute Program",
    "Waste / Garbage": "Elevator / Chute Program",
    "Terrace / Outdoor": "Outdoor / Patio",
    "Washroom / Change Room": "Washrooms",
    "Office / Admin": "Office / Admin",
    "Storage": "Back-of-House",
    "Mechanical / Back-of-House": "Back-of-House",
    "Loading Dock": "Loading Dock",
}
AMENITY_AREAS = {"Amenity", "Fitness", "Pool / Spa"}
# Deterministic tie-break order when two area types are equally common in a cell.
AREA_PRIORITY = ["Lobby", "Amenity", "Fitness", "Pool / Spa", "Guest Suite", "Parking",
                 "Terrace / Outdoor", "Washroom / Change Room", "Corridor / Hallway",
                 "Elevator", "Waste / Garbage", "Stairwell", "Loading Dock",
                 "Office / Admin", "Storage", "Mechanical / Back-of-House", "Other / Unknown"]
LOBBY_KW = ("lobby", "enterphone", "entrance", "security", "cafe", "mail",
            "reception", "concierge", "front")
STAFF_KW = ("lunch room", "lunchroom", "break room", "breakroom", "fridge",
            "staff room", "kitchenette", "locker")
SPECIAL_KW = ("special", "detail", "high reach", "scrub", "steam", "squeegee",
              "baseboard", "spot clean", "window", "glass", "carpet", "polish", "wax",
              "re-sweep", "resweep", "re-clean", "reclean", "re-check", "recheck",
              "re-mop", "deep clean", "shampoo", "touch up", "touch-up")


def resolve_source_csv() -> str:
    norm = os.path.join(REPO, "schedule_data", "normalized", "Parker_cleaning_responsibilities.csv")
    return norm if os.path.exists(norm) else os.path.join(
        REPO, "schedule_data", "Parker_cleaning_responsibilities.csv")


def resolve_input_workbook() -> str:
    hits = [p for p in glob.glob(os.path.join(REPO, "**", "Parker_Operations_Workbook.xlsx"),
                                 recursive=True) if ".git" not in p]
    if not hits:
        raise FileNotFoundError("Parker_Operations_Workbook.xlsx not found")
    return hits[0]


def expand_days(v: str) -> list[str]:
    low = re.sub(r"^[\s\-–•*]+", "", (v or "")).strip().lower()
    low = re.sub(r"\s+", " ", low)
    if low in FULL:
        return [WEEK[FULL[low]]]
    if "-" in low:
        a, _, b = low.partition("-")
        if a.strip() in ABBR and b.strip() in ABBR:
            i, j = ABBR[a.strip()], ABBR[b.strip()]
            idx = range(i, j + 1) if i <= j else list(range(i, 7)) + list(range(0, j + 1))
            return [WEEK[k] for k in idx]
    return [WEEK[ABBR[low]]] if low in ABBR else []


def start_min(t: str) -> int:
    m = re.match(r"^\s*(\d{1,2}):(\d{2})\s*([AP]M)?\s*$", (t or "").strip(), re.I)
    if not m:
        return 10 ** 9
    h = int(m.group(1)) % 12
    if (m.group(3) or "").upper() == "PM":
        h += 12
    return h * 60 + int(m.group(2))


def floor_of(loc: str) -> str:
    m = re.search(r"(\d{1,2})(st|nd|rd|th)\s*floor", loc.lower())
    if m:
        return f"{m.group(1)}{m.group(2)} Floor"
    if re.search(r"\buph\b|upper penthouse", loc.lower()):
        return "Upper PH"
    return ""


def cell_label(items: list[dict]) -> str:
    """Summarise one employee/day/time-block's assignments into a single label."""
    if not items:
        return ""
    texts = " ".join((i["loc"] + " " + i["task"]).lower() for i in items)
    # Lunch/Break only for an actual break — not a "Lunch Room" / "Break Room" to clean.
    if re.search(r"\blunch\b", texts) and "lunch room" not in texts and "lunchroom" not in texts:
        return "Lunch"
    if re.search(r"\bbreak\b", texts) and "break room" not in texts and "breakroom" not in texts:
        return "Break"

    # Multi-area ground-floor sweep -> "Lobby Cluster"
    lobby_hits = sum(1 for i in items if any(k in i["loc"].lower() for k in LOBBY_KW))
    if len(items) >= 3 and lobby_hits >= 2 and ("lobby" in texts or "enterphone" in texts):
        return "Lobby Cluster"

    areas = Counter(i["area"] for i in items)
    mx = max(areas.values())
    leaders = [a for a in AREA_PRIORITY if areas.get(a, 0) == mx]
    dom = leaders[0] if leaders else sorted(a for a, c in areas.items() if c == mx)[0]

    if dom in AMENITY_AREAS:
        floors = sorted({floor_of(i["loc"]) for i in items
                         if i["area"] in AMENITY_AREAS and floor_of(i["loc"])})
        if len(floors) == 1:      # only prefix a floor when it's unambiguous
            return f"{floors[0]} Amenities"
        return "Amenities"
    if dom == "Other / Unknown":
        if any(k in texts for k in STAFF_KW):
            return "Back-of-House"
        if any(k in texts for k in SPECIAL_KW):
            return "Specials"
        if "debris" in texts or "trash can" in texts:
            return "Outdoor / Patio"
        return "Review Needed"
    return AREA_LABEL.get(dom, "Review Needed")


def build_grid():
    with open(resolve_source_csv(), newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    cells: dict[tuple, list] = defaultdict(list)   # (day, block, employee) -> items
    block_start: dict[str, int] = {}
    employees: set[str] = set()
    for r in rows:
        emp = clean(r.get("employee_name", ""))
        if not emp:
            continue
        employees.add(emp)
        s, e = clean(r.get("start_time", "")), clean(r.get("end_time", ""))
        block = f"{s}–{e}" if s and e else (s or e)
        block_start[block] = start_min(s)
        loc = clean(r.get("location", ""))
        item = {"loc": loc, "area": classify_area(clean_location(loc)),
                "task": clean(r.get("task", ""))}
        for day in expand_days(clean(r.get("day_of_week", ""))):
            cells[(day, block, emp)].append(item)

    employees = sorted(employees)
    row_keys = sorted({(d, b) for (d, b, _) in cells},
                      key=lambda k: (WEEK.index(k[0]), block_start.get(k[1], 0), k[1]))

    grid = []
    review = []
    for day, block in row_keys:
        cellvals = []
        for emp in employees:
            label = cell_label(cells.get((day, block, emp), []))
            cellvals.append(label)
            if label == "Review Needed":
                review.append((day, block, emp))
        grid.append([day, block] + cellvals)
    return employees, grid, review, cells


def main():
    employees, grid, review, _cells = build_grid()

    in_wb = resolve_input_workbook()
    wb = load_workbook(in_wb)
    if "Schedule Grid" in wb.sheetnames:
        del wb["Schedule Grid"]
    ws = wb.create_sheet("Schedule Grid", 0)   # first sheet
    ws.sheet_view.showGridLines = False

    ncols = 2 + len(employees)
    head_fill = PatternFill("solid", fgColor="1A1A1A")
    head_font = Font(bold=True, color="FFFFFF")
    band = PatternFill("solid", fgColor="F4F4F4")
    review_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    head_align = Alignment(wrap_text=True, vertical="center", horizontal="center")

    # note row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    n = ws.cell(1, 1, NOTE); n.font = Font(italic=True, size=9, color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 26

    # header row
    headers = ["Day", "Time Block"] + employees
    for j, h in enumerate(headers, start=1):
        c = ws.cell(2, j, h); c.fill = head_fill; c.font = head_font
        c.alignment = head_align; c.border = border
    ws.row_dimensions[2].height = 20

    # data rows (light band alternating by day)
    day_order = {d: i for i, d in enumerate(WEEK)}
    for i, row in enumerate(grid):
        r = i + 3
        shade = day_order[row[0]] % 2 == 1
        for j, val in enumerate(row, start=1):
            c = ws.cell(r, j, val); c.border = border; c.alignment = wrap
            if j <= 2:
                c.font = Font(bold=(j == 1))
            if val == "Review Needed":
                c.fill = review_fill
            elif shade:
                c.fill = band

    widths = [13, 18] + [20] * len(employees)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws.freeze_panes = "C3"  # freeze note+header rows and Day/Time Block columns
    last = len(grid) + 2
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{last}"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    try:
        wb.save(OUT)
    except PermissionError:
        print(f"WARNING: {os.path.basename(OUT)} is open — close it and rerun.")
        return

    print("output path:", os.path.relpath(OUT, REPO))
    print("source workbook (unchanged):", os.path.relpath(in_wb, REPO))
    print("employees included:", employees)
    print("grid rows (Day x Time Block):", len(grid))
    print("cells marked 'Review Needed':", len(review))
    for d, b, e in review[:20]:
        print(f"   {d} {b} — {e}")


if __name__ == "__main__":
    main()
